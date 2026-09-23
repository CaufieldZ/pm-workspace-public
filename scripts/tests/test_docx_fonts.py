"""docx_fonts 单测：生成器产出的 docx 不该带日语主题语言。

锁住三件事：
  检测  — python-docx 默认模板带 ja-JP（本文件的存在理由），inspect_docx 能读出来
  归正  — new_document() 建出来的文档是 zh-CN；normalize_docx() 能修已落盘的旧产物
  幂等  — 已经是 zh-CN 的文档再跑一遍：不改字节、不改文件
"""
from pathlib import Path

import pytest
from docx import Document

from lib.docx_fonts import DEFAULT_LANG, ensure_ea_language, inspect_docx, new_document, normalize_docx


def _save_default(path: Path) -> Path:
    """用 python-docx 默认模板存一份（继承上游默认，用来当「带病样本」）。"""
    doc = Document()
    doc.add_paragraph("中文测试")
    doc.save(str(path))
    return path


# ── 检测 ────────────────────────────────────────────────
def test_python_docx_default_is_japanese(tmp_path):
    """python-docx 默认模板的主题语言是日语——这正是要归正的对象。"""
    p = _save_default(tmp_path / "default.docx")
    info = inspect_docx(p)
    assert info["theme_lang"] == "ja-JP"
    assert info["has_japanese_lang"] is True


# ── 归正：内存（生成时） ────────────────────────────────
def test_new_document_is_chinese(tmp_path):
    """new_document() 建出来的文档，主题语言已是 zh-CN。"""
    p = tmp_path / "new.docx"
    doc = new_document()
    doc.add_paragraph("中文测试")
    doc.save(str(p))
    info = inspect_docx(p)
    assert info["theme_lang"] == DEFAULT_LANG
    assert info["has_japanese_lang"] is False


def test_ensure_ea_language_on_plain_document(tmp_path):
    """对已有 Document 就地归正（生成器第二种接法）。"""
    p = tmp_path / "inplace.docx"
    doc = Document()
    ensure_ea_language(doc)
    doc.add_paragraph("中文测试")
    doc.save(str(p))
    assert inspect_docx(p)["theme_lang"] == DEFAULT_LANG


# ── 归正：落盘（收尾校） ────────────────────────────────
def test_normalize_patches_saved_document(tmp_path):
    """旧产物（已带 ja-JP）能被 normalize_docx 修好。"""
    p = _save_default(tmp_path / "legacy.docx")
    assert inspect_docx(p)["has_japanese_lang"] is True
    report = normalize_docx(p)
    assert report["changed"] is True
    assert report["before"]["theme_lang"] == "ja-JP"
    assert report["after"]["theme_lang"] == DEFAULT_LANG
    assert inspect_docx(p)["has_japanese_lang"] is False


def test_normalize_sets_theme_east_asian_font(tmp_path):
    """传 ea_font 时，主题东亚字体与简体中文脚本一并指向它。"""
    p = _save_default(tmp_path / "font.docx")
    normalize_docx(p, ea_font="Microsoft YaHei")
    info = inspect_docx(p)
    assert 'typeface="Microsoft YaHei"' in str(info["theme_ea"])
    assert 'typeface="Microsoft YaHei"' in str(info["theme_hans"])


def test_normalize_is_idempotent(tmp_path):
    """已归正的文档再跑：report.changed 为假，文件字节不变。"""
    p = _save_default(tmp_path / "twice.docx")
    normalize_docx(p)
    snapshot = p.read_bytes()
    report = normalize_docx(p)
    assert report["changed"] is False
    assert p.read_bytes() == snapshot


if __name__ == "__main__":
    pytest.main([__file__, "-q"])
